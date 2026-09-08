"""
Tests for modules/resume_docx.py — the .docx renderer and its structural
inspector. No API calls anywhere here.

Run: python3 -m unittest discover -s tests -v
"""
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

try:
    from docx import Document
    from modules import resume_docx
    _HAS_DOCX = True
except ImportError:  # python-docx only installed in .venv
    _HAS_DOCX = False

from modules import resume_style as S

RESUME = {
    "summary": "Program leader focused on data platforms.",
    "experience": [
        {"title": "Senior Technical Program Manager", "company": "INRIX",
         "location": "Kirkland, WA", "start": "Jun 2022", "end": "Nov 2024",
         "current": False,
         "bullets": ["Ran delivery across eight teams and hit 99% on-time delivery.",
                     "Brought platform cost down by 40 percent."]},
        {"title": "Technical Program Manager", "company": "IHME",
         "location": "Seattle, WA", "start": "2019", "end": "2022", "current": False,
         "bullets": ["Raised accessibility from 30 to 89 percent with clinician research."]},
    ],
    "skills": ["Program management", "SQL", "AWS"],
    "education": [{"credential": "B.S. Computer Science", "institution": "University of Washington", "year": "2017"}],
    "certifications": [{"name": "PMP", "year": None}, {"name": "Certified Scrum Master", "year": "2020"}],
}
CONTACT = {"name": "Jane Doe", "email": "jane@example.com", "phone": "(206) 555-0100",
           "location": "Seattle, WA", "linkedin": "linkedin.com/in/janedoe",
           "github": "", "website": ""}

# Populated github/website, to prove they're excluded by default rather than
# only "happening" to be absent because CONTACT above leaves them blank.
CONTACT_WITH_LINKS = dict(CONTACT, github="github.com/janedoe/some-repo",
                          website="https://janedoe-portfolio.example.com")


@unittest.skipUnless(_HAS_DOCX, "python-docx not installed for this interpreter (.venv only)")
class TestRenderProducesReadableDocx(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()
        self.path = os.path.join(self.d, "r.docx")

    def test_no_summary_by_default(self):
        resume_docx.render(RESUME, CONTACT, self.path, include_summary=False)
        self.assertTrue(resume_docx.opens_cleanly(self.path))
        texts = [p.text.strip().upper() for p in Document(self.path).paragraphs]
        self.assertNotIn("SUMMARY", texts)
        self.assertIn("EXPERIENCE", texts)

    def test_summary_included_when_asked(self):
        resume_docx.render(RESUME, CONTACT, self.path, include_summary=True)
        texts = [p.text.strip().upper() for p in Document(self.path).paragraphs]
        self.assertIn("SUMMARY", texts)

    def test_structure_is_ats_safe(self):
        resume_docx.render(RESUME, CONTACT, self.path, include_summary=False)
        s = resume_docx.inspect_structure(self.path)
        self.assertEqual(s["tables"], 0)
        self.assertEqual(s["images"], 0)
        self.assertEqual(s["text_boxes"], 0)
        self.assertFalse(s["has_header_content"])
        self.assertFalse(s["has_footer_content"])
        self.assertEqual(s["columns"], 1)

    def test_section_order_no_summary(self):
        resume_docx.render(RESUME, CONTACT, self.path, include_summary=False)
        s = resume_docx.inspect_structure(self.path)
        got = [h for h in s["section_headings"] if h in set(S.SECTION_ORDER_NO_SUMMARY)]
        self.assertEqual(got, S.SECTION_ORDER_NO_SUMMARY)

    def test_section_order_with_summary(self):
        resume_docx.render(RESUME, CONTACT, self.path, include_summary=True)
        s = resume_docx.inspect_structure(self.path)
        got = [h for h in s["section_headings"] if h in set(S.SECTION_ORDER_WITH_SUMMARY)]
        self.assertEqual(got, S.SECTION_ORDER_WITH_SUMMARY)

    def test_margins_and_single_column_from_style_reference(self):
        resume_docx.render(RESUME, CONTACT, self.path, include_summary=False)
        doc = Document(self.path)
        sec = doc.sections[0]
        self.assertAlmostEqual(sec.left_margin.inches, S.MARGIN_IN, places=2)
        self.assertAlmostEqual(sec.top_margin.inches, S.MARGIN_IN, places=2)

    def test_contact_details_present_in_body(self):
        resume_docx.render(RESUME, CONTACT, self.path, include_summary=False)
        full = "\n".join(p.text for p in Document(self.path).paragraphs)
        self.assertIn("Jane Doe", full)
        self.assertIn("jane@example.com", full)
        self.assertIn("Seattle, WA", full)

    def test_no_em_dash_in_output_and_ranges_use_hyphens(self):
        # Em dash is never produced. En dash is *permitted* in ranges by the
        # validator, but the renderer deterministically uses a plain hyphen, so
        # the emitted document happens to contain neither dash character.
        resume_docx.render(RESUME, CONTACT, self.path, include_summary=True)
        full = "\n".join(p.text for p in Document(self.path).paragraphs)
        self.assertNotIn("—", full)                      # em dash: hard rule
        self.assertIn("Jun 2022 - Nov 2024", full)       # range joined with a hyphen

    def test_no_contact_heading_and_name_is_first_paragraph(self):
        # 2026-09-05: a labeled "Contact" heading above the name/contact line
        # read as filler on a one-page resume. Name now leads the document
        # directly, with no heading paragraph above it.
        resume_docx.render(RESUME, CONTACT, self.path, include_summary=False)
        paragraphs = [p for p in Document(self.path).paragraphs if p.text.strip()]
        texts_upper = [p.text.strip().upper() for p in paragraphs]
        self.assertNotIn("CONTACT", texts_upper)
        self.assertEqual(paragraphs[0].text.strip(), CONTACT["name"])

    def test_github_and_website_excluded_from_contact_line_by_default(self):
        # Job-relevant links only by default -- the job-search-agent repo and
        # personal art-website links must not appear unless explicitly
        # requested via contact_fields, even when populated on the contact dict.
        resume_docx.render(RESUME, CONTACT_WITH_LINKS, self.path, include_summary=False)
        full = "\n".join(p.text for p in Document(self.path).paragraphs)
        self.assertNotIn(CONTACT_WITH_LINKS["github"], full)
        self.assertNotIn(CONTACT_WITH_LINKS["website"], full)
        self.assertIn(CONTACT_WITH_LINKS["linkedin"], full)   # job-relevant link still present

    def test_github_and_website_included_when_explicitly_requested(self):
        resume_docx.render(RESUME, CONTACT_WITH_LINKS, self.path, include_summary=False,
                            contact_fields=("location", "phone", "email", "linkedin",
                                            "github", "website"))
        full = "\n".join(p.text for p in Document(self.path).paragraphs)
        self.assertIn(CONTACT_WITH_LINKS["github"], full)
        self.assertIn(CONTACT_WITH_LINKS["website"], full)

    def test_role_header_and_date_line_keep_with_next(self):
        # Chains heading -> date line -> first bullet so a page break can
        # never strand the heading alone with its bullets pushed to the next
        # page. Only header + date-subline need the flag; the first bullet
        # inherits the "stay with the paragraph before it" effect from the
        # date-subline's own keep_with_next.
        resume_docx.render(RESUME, CONTACT, self.path, include_summary=False)
        paragraphs = Document(self.path).paragraphs
        header_p = next(p for p in paragraphs if p.text.strip() == "Senior Technical Program Manager, INRIX, Kirkland, WA")
        idx = paragraphs.index(header_p)
        date_p = paragraphs[idx + 1]
        self.assertTrue(header_p.paragraph_format.keep_with_next)
        self.assertTrue(date_p.paragraph_format.keep_with_next)

    def test_bullets_do_not_force_keep_with_next(self):
        # Scoped narrowly to heading+first-bullet by request -- bullets
        # themselves are not chained together, so a role with many bullets
        # can still flow/split across a page instead of being forced onto
        # the next page as one unbreakable block.
        resume_docx.render(RESUME, CONTACT, self.path, include_summary=False)
        bullet_p = next(p for p in Document(self.path).paragraphs
                        if p.text.strip().startswith(S.BULLET_CHAR))
        self.assertFalse(bullet_p.paragraph_format.keep_with_next)


if __name__ == "__main__":
    unittest.main()
